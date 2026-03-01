"""
Convert Markdown to Word (.docx) using python-docx.
Run: pip install python-docx
     python md_to_docx.py
"""
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Installing python-docx...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_paragraph_with_formatting(doc, line: str) -> None:
    """Add a paragraph, preserving **bold** and leaving [refs] as-is."""
    p = doc.add_paragraph()
    rest = line.strip()
    while rest:
        m = re.search(r"\*\*(.+?)\*\*", rest)
        if m:
            before = rest[: m.start()]
            if before:
                p.add_run(before)
            r = p.add_run(m.group(1))
            r.bold = True
            rest = rest[m.end() :]
        else:
            p.add_run(rest)
            break


def md_to_docx(md_path: str, docx_path: str) -> None:
    md_path = Path(md_path)
    docx_path = Path(docx_path)
    text = md_path.read_text(encoding="utf-8", errors="replace")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.name = "Times New Roman"

    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("---"):
            pass
        elif line.strip() == "":
            doc.add_paragraph()
        elif line.strip().startswith("[Figure ") or line.strip().startswith("[Table "):
            p = doc.add_paragraph()
            r = p.add_run(line.strip())
            r.italic = True
        elif line.strip().startswith("**") and line.strip().endswith("**") and line.strip().count("**") == 2:
            p = doc.add_paragraph()
            p.add_run(line.strip().strip("* ")).bold = True
        elif line.strip():
            add_paragraph_with_formatting(doc, line)
        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    md_file = Path(__file__).parent / "FC6P01ES_Interim_Report_AAC_Autism.md"
    docx_file = Path(__file__).parent / "FC6P01ES_Interim_Report_AAC_Autism.docx"
    if not md_file.exists():
        print(f"Not found: {md_file}")
        sys.exit(1)
    md_to_docx(str(md_file), str(docx_file))
