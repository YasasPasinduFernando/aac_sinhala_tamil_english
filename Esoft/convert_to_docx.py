"""
Convert FC6P01ES_Interim_Report_FINAL.md to a Word document (.docx)
Times New Roman, size 12, 1.5 line spacing, black headings, simple academic formatting.
"""

import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

INPUT_MD = os.path.join(os.path.dirname(__file__), "FC6P01ES_Interim_Report_FINAL.md")
OUTPUT_DOCX = os.path.join(os.path.dirname(__file__), "FC6P01ES_Interim_Report_FINAL.docx")

FONT = 'Times New Roman'
FONT_SIZE = 12
LINE_SPACING = 1.5
BLACK = RGBColor(0, 0, 0)


def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ['top', 'left', 'bottom', 'right']:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_run(paragraph, text, bold=False, italic=False, size=FONT_SIZE):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic
    return run


def process_inline(paragraph, text, size=FONT_SIZE, base_bold=False):
    """Handle **bold** and *italic* markers within text."""
    pattern = r'(\*\*.*?\*\*|\*.*?\*|[^*]+)'
    parts = re.findall(pattern, text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            add_run(paragraph, part[2:-2], bold=True, size=size)
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            add_run(paragraph, part[1:-1], italic=True, size=size)
        else:
            add_run(paragraph, part, bold=base_bold, size=size)


def set_paragraph_spacing(p, before=0, after=6):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = LINE_SPACING


def parse_table(lines, start_idx):
    rows = []
    i = start_idx
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('|') and line.endswith('|'):
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if cells and not all(re.match(r'^[-:]+$', c) for c in cells):
                rows.append(cells)
            i += 1
        else:
            break
    return rows, i


def convert(input_path, output_path):
    with open(input_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # Default style
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(FONT_SIZE)
    style.font.color.rgb = BLACK
    style.paragraph_format.line_spacing = LINE_SPACING
    style.paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        # Skip blanks
        if stripped == '':
            i += 1
            continue

        # Skip horizontal rules
        if stripped == '---':
            i += 1
            continue

        # H1 — Document title (Heading 1)
        if stripped.startswith('# ') and not stripped.startswith('## '):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style='Heading 1')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=24, after=12)
            # Ensure font is consistent with ESOFT requirements
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(16)
                run.font.color.rgb = BLACK
            if not p.runs:
                add_run(p, text, bold=True, size=16)
            else:
                p.runs[0].text = text
            i += 1
            continue

        # H2 — Major section (Heading 2)
        if stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph(style='Heading 2')
            set_paragraph_spacing(p, before=18, after=8)
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(14)
                run.font.color.rgb = BLACK
            if not p.runs:
                add_run(p, text, bold=True, size=14)
            else:
                p.runs[0].text = text
            i += 1
            continue

        # H3 — Subsection (Heading 3)
        if stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph(style='Heading 3')
            set_paragraph_spacing(p, before=14, after=6)
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(12)
                run.font.color.rgb = BLACK
            if not p.runs:
                add_run(p, text, bold=True, size=12)
            else:
                p.runs[0].text = text
            i += 1
            continue

        # H4 — Sub-subsection (Heading 4)
        if stripped.startswith('#### '):
            text = stripped[5:].strip()
            p = doc.add_paragraph(style='Heading 4')
            set_paragraph_spacing(p, before=10, after=4)
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(12)
                run.font.color.rgb = BLACK
            if not p.runs:
                add_run(p, text, bold=True, italic=True, size=12)
            else:
                p.runs[0].text = text
            i += 1
            continue

        # Table caption [Table N: ...]
        if re.match(r'^\[Table \d+:.*\]$', stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=8, after=4)
            add_run(p, stripped[1:-1], bold=True, italic=True, size=10)
            i += 1
            continue

        # Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            rows, next_i = parse_table(lines, i)
            if rows:
                num_cols = len(rows[0])
                table = doc.add_table(rows=len(rows), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                for row_idx, row_data in enumerate(rows):
                    for col_idx, cell_text in enumerate(row_data):
                        if col_idx < num_cols:
                            cell = table.cell(row_idx, col_idx)
                            cell.text = ''
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.0

                            is_header = (row_idx == 0)
                            process_inline(p, cell_text, size=10, base_bold=is_header)
                            set_cell_borders(cell)

                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i = next_i
            continue

        # Cover page fields: **Key:** Value
        if stripped.startswith('**') and ':**' in stripped:
            p = doc.add_paragraph()
            set_paragraph_spacing(p)
            match = re.match(r'\*\*(.*?):\*\*\s*(.*)', stripped)
            if match:
                add_run(p, f"{match.group(1)}: ", bold=True)
                if match.group(2):
                    add_run(p, match.group(2))
                if match.group(1) in ['Module Code', 'Module Name', 'Institution',
                                      'Project Title', 'Student Name',
                                      'Academic Year', 'Submission Date']:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                process_inline(p, stripped)
            i += 1
            continue

        # Keywords
        if stripped.startswith('**Keywords:**'):
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=8)
            add_run(p, 'Keywords: ', bold=True, size=11)
            add_run(p, stripped.replace('**Keywords:**', '').strip(), italic=True, size=11)
            i += 1
            continue

        # Signature lines
        if stripped.startswith('Signed:') or stripped.startswith('Date:'):
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=14)
            add_run(p, stripped)
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match and not stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            set_paragraph_spacing(p, after=3)
            add_run(p, f"{num_match.group(1)}. ", bold=True)
            process_inline(p, num_match.group(2))
            i += 1
            continue

        # Bullet list
        if stripped.startswith('- '):
            text = stripped[2:]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.27)
            p.paragraph_format.first_line_indent = Cm(-0.63)
            set_paragraph_spacing(p, after=3)
            add_run(p, '\u2022 ')
            process_inline(p, text)
            i += 1
            continue

        # Research question labels
        rq_match = re.match(r'^\*\*(RQ\d+):\*\*\s*(.*)', stripped)
        if rq_match:
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=8, after=2)
            add_run(p, f"{rq_match.group(1)}: ", bold=True)
            add_run(p, rq_match.group(2), italic=True)
            i += 1
            continue

        # Increment/Phase headings
        inc_match = re.match(r'^\*\*((?:Increment|Phase) \d+.*?)\*\*\s*(.*)', stripped)
        if inc_match:
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=8, after=3)
            add_run(p, inc_match.group(1) + ' ', bold=True)
            if inc_match.group(2):
                add_run(p, inc_match.group(2))
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        set_paragraph_spacing(p)
        process_inline(p, stripped)
        i += 1

    doc.save(output_path)
    print(f"Saved: {output_path}")
    print(f"Size: {os.path.getsize(output_path) / 1024:.1f} KB")


if __name__ == '__main__':
    convert(INPUT_MD, OUTPUT_DOCX)
